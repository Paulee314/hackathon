"""The offers sheet for the demo: every offer in the month, each rule's verdict, and the rules compared.

    .venv/bin/python -m demo.offers_sheet samples/october     # writes demo/ride_offers_october.xlsx

The verdicts come from the replay in car/rules.py; profits on the sheet are live formulas off an Assumptions
tab, so the totals recalculate when an input changes. Needs openpyxl (requirements.txt).
"""
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from car.rules import export

HERE = Path(__file__).parent
MONTH = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE.parent / "samples" / "october"
_x = export(MONTH)
_names = [p["name"] for p in _x["presets"]]
_L = {"take": "T", "skip": "S", "busy": "B"}
d = {"names": _names, "totals": _x["totals"], "hours": _x["hours_on_road"], "tip": _x["tip_rate_exact"],
     "kwh_mi": _x["kwh_per_mile"], "profile": _x["profile"], "queue": _x["queue_minutes"], "off": _x["off_after_minutes"],
     "rows": [[r["at"][5:16], r["platform"][0], r["pickup_miles"], r["pickup_minutes"], r["trip_miles"], r["trip_minutes"],
               r["fare"], int(r["accepted"]), r["actual_fare"], r["actual_profit"],
               "".join(_L[r["decisions"][n]] for n in _names), r["id"][:6]] for r in _x["rows"]]}
rows = d["rows"]
N = len(rows)
FIRST, LAST = 2, N + 1

FONT = "Arial"
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=14, bold=True, color="2D2A5A")
f_note = Font(name=FONT, size=9, italic=True, color="5C5A86")
f_input = Font(name=FONT, size=10, color="0000FF")      # hardcoded inputs
f_link = Font(name=FONT, size=10, color="008000")       # links to another sheet
fill_head = PatternFill("solid", fgColor="2D2A5A")
fill_band = PatternFill("solid", fgColor="EEF6FF")
fill_key = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="DCEBFF")
MONEY = '$#,##0.00;($#,##0.00);"-"'

wb = Workbook()

# ---------------- Assumptions ----------------
A = wb.active
A.title = "Assumptions"
p = d["profile"]
A["A1"] = "Assumptions behind the estimates"
A["A1"].font = f_title
A.append([])
A.append(["Input", "Value", "Where it comes from"])
assumptions = [
    ("Electricity used per mile (kWh)", d["kwh_mi"], "0.000", "Tessie drives for the month: energy used over miles driven"),
    ("Home electricity ($/kWh)", p["home_electricity_per_kwh"], "$0.00", "Owner's vehicle profile"),
    ("Supercharger ($/kWh)", p["supercharger_per_kwh"], "$0.00", "Owner's vehicle profile"),
    ("Share of charging at Superchargers", p["share_of_energy_from_supercharging"], "0%", "Owner's vehicle profile"),
    ("Blended electricity ($/kWh)", "=B5*(1-B7)+B6*B7", "$0.000", "Home and Supercharger rates, weighted by share"),
    ("Wear on the car ($/mile)", p["depreciation_per_mile"], "$0.00", "Owner's vehicle profile (depreciation)"),
    ("Tire set ($)", p["tire_set_cost"], "$#,##0", "Owner's vehicle profile"),
    ("Tire life (miles)", p["tire_life_miles"], "#,##0", "Owner's vehicle profile"),
    ("Tire wear ($/mile)", "=B10/B11", "$0.0000", "Tire set over tire life"),
    ("Car cost per mile ($)", "=B4*B8+B9+B12", "$0.0000", "Electricity plus wear plus tires"),
    ("Tip rate on rides driven", d["tip"], "0.0%", "Rides driven this month: tips as a share of the fares on their offer cards"),
    ("Hours on the road", d["hours"], "0.0", f"Real offers and rides, in stretches; a gap over {d['off']} minutes counts as off"),
    ("Queue window (minutes)", d["queue"], "0", "A passed-on offer counts only if the car is free within this long"),
]
for i, (label, value, fmt, src) in enumerate(assumptions, start=4):
    A.cell(row=i, column=1, value=label).font = f_base
    c = A.cell(row=i, column=2, value=value)
    c.number_format = fmt
    c.font = f_base if isinstance(value, str) and value.startswith("=") else f_input
    A.cell(row=i, column=3, value=src).font = f_note
for col in "ABC":
    A[f"{col}3"].font = f_head
    A[f"{col}3"].fill = fill_head
for r in (13, 14):
    A[f"B{r}"].fill = fill_key
A["A18"] = ("Blue values are inputs you can change. Changing them updates the estimated profit of offers that "
            "were passed on, and every total on the Rules tab.")
A["A19"] = ("Which offers each rule takes (Take, Skip, Busy on the Offers tab) comes from the replay in car/rules.py "
            "and does not change when you edit these inputs.")
A["A20"] = "Synthetic data for a fictional operator, Bayou City Mobility LLC. Advice only: no app accepts or declines rides for you."
for r in (18, 19, 20):
    A[f"A{r}"].font = f_note
A.column_dimensions["A"].width = 36
A.column_dimensions["B"].width = 14
A.column_dimensions["C"].width = 90

COST = "Assumptions!$B$13"
TIP = "Assumptions!$B$14"
HOURS = "Assumptions!$B$15"

# ---------------- Offers ----------------
O = wb.create_sheet("Offers")
rule_names = d["names"][1:]
headers = ["Offered", "Platform", "Pickup (mi)", "Pickup (min)", "Trip (mi)", "Trip (min)", "Fare on card ($)",
           "Fare per mile ($)", "Card $/hr after car costs", "You", "Fare + tip, actual ($)", "True profit, actual ($)",
           "True profit used ($)"] + rule_names + ["Offer ID"]
O.append(headers)
VERDICT = {"T": "Take", "S": "Skip", "B": "Busy"}
for i, r in enumerate(rows, start=FIRST):
    when = datetime.strptime("2026-" + r[0], "%Y-%m-%d %H:%M")
    took = bool(r[7])
    O.cell(row=i, column=1, value=when).number_format = "ddd mmm d, h:mm AM/PM"
    O.cell(row=i, column=2, value="Uber" if r[1] == "u" else "Lyft")
    for col, v in zip(range(3, 8), r[2:7]):
        O.cell(row=i, column=col, value=v).font = f_input
    O.cell(row=i, column=8, value=f"=IFERROR(G{i}/(C{i}+E{i}),0)")
    O.cell(row=i, column=9, value=f"=IFERROR((G{i}-(C{i}+E{i})*{COST})/((D{i}+F{i})/60),0)")
    O.cell(row=i, column=10, value="Took" if took else "Passed")
    if took:
        O.cell(row=i, column=11, value=r[8]).font = f_input
        O.cell(row=i, column=12, value=r[9]).font = f_input
    O.cell(row=i, column=13, value=f'=IF(J{i}="Took",L{i},ROUND(G{i}*(1+{TIP})-(C{i}+E{i})*{COST},2))')
    for k, letter in enumerate(r[10][1:]):
        O.cell(row=i, column=14 + k, value=VERDICT[letter])
    O.cell(row=i, column=14 + len(rule_names), value=r[11])
    for col in range(1, len(headers) + 1):
        c = O.cell(row=i, column=col)
        if c.font != f_input:
            c.font = f_base
        if i % 2 == 0:
            c.fill = fill_band
for col in range(1, len(headers) + 1):
    c = O.cell(row=1, column=col)
    c.font = f_head
    c.fill = fill_head
    c.alignment = Alignment(wrap_text=True, vertical="center")
O.row_dimensions[1].height = 44
fmt = {3: "0.00", 4: "0.0", 5: "0.00", 6: "0.0", 7: MONEY, 8: MONEY, 9: MONEY, 11: MONEY, 12: MONEY, 13: MONEY}
for col, f in fmt.items():
    for i in range(FIRST, LAST + 1):
        O.cell(row=i, column=col).number_format = f
widths = {1: 20, 2: 9, 3: 9, 4: 9, 5: 9, 6: 9, 7: 11, 8: 10, 9: 12, 10: 9, 11: 12, 12: 12, 13: 12}
for col, w in widths.items():
    O.column_dimensions[get_column_letter(col)].width = w
for k in range(len(rule_names)):
    O.column_dimensions[get_column_letter(14 + k)].width = 13
O.column_dimensions[get_column_letter(14 + len(rule_names))].width = 10
O.freeze_panes = "B2"
O.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{LAST}"
rng = f"{get_column_letter(14)}{FIRST}:{get_column_letter(13 + len(rule_names))}{LAST}"
O.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Take"'], fill=PatternFill("solid", fgColor="D6F5E3")))
O.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Skip"'], fill=PatternFill("solid", fgColor="FFE3F1")))
O.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Busy"'], fill=PatternFill("solid", fgColor="E6E6EE")))
O["J1"].comment = Comment("Took: the driver accepted the offer and drove it. Passed: the driver declined it.", "hackathon")
O["L1"].comment = Comment("From the ride's own drives: fare and tip minus electricity, wear and tires, with the pickup "
                          "miles charged to the ride (car/profit.py).", "hackathon")
O["M1"].comment = Comment("Actual true profit for rides driven; for offers passed on, an estimate from the offer card: "
                          "fare plus the month's tip rate, minus card miles at the car cost per mile.", "hackathon")
O["N1"].comment = Comment("Take: the rule takes it. Skip: the rule turns it down. Busy: the rule would take it, but the car "
                          "is on another ride. Computed by the replay in car/rules.py.", "hackathon")

# ---------------- Rules (summary, first tab) ----------------
R = wb.create_sheet("Rules", 0)
R["A1"] = "Ride rules, replayed on every offer this month"
R["A1"].font = f_title
R["A2"] = ("Each rule sits on top of what you did: it drops the rides you took that it would have skipped, then fills the "
           "freed time with offers you passed on that fit. True profit is after electricity, wear on the car and tires.")
R["A2"].font = f_note
R["A2"].alignment = Alignment(wrap_text=True, vertical="top")
R.merge_cells("A2:G2")
R.row_dimensions[2].height = 30
heads = ["Rule", "Rides", "True profit ($)", "Per hour on the road ($)", "Versus what you did ($)",
         "Rides you took that it skips", "Offers you passed that it takes"]
R.append([])
R.append(heads)
for col in range(1, len(heads) + 1):
    c = R.cell(row=4, column=col)
    c.font = f_head
    c.fill = fill_head
    c.alignment = Alignment(wrap_text=True, vertical="center")
R.row_dimensions[4].height = 32
J = f"Offers!$J${FIRST}:$J${LAST}"
M = f"Offers!$M${FIRST}:$M${LAST}"
R.cell(row=5, column=1, value="What you did")
R.cell(row=5, column=2, value=f'=COUNTIF({J},"Took")')
R.cell(row=5, column=3, value=f'=SUMIFS({M},{J},"Took")')
R.cell(row=5, column=4, value=f"=C5/{HOURS}")
R.cell(row=5, column=5, value="=C5-$C$5")
R.cell(row=5, column=6, value=0)
R.cell(row=5, column=7, value=0)
for k, name in enumerate(rule_names):
    row = 6 + k
    col = get_column_letter(14 + k)
    V = f"Offers!${col}${FIRST}:${col}${LAST}"
    R.cell(row=row, column=1, value=name)
    R.cell(row=row, column=2, value=f'=COUNTIF({V},"Take")')
    R.cell(row=row, column=3, value=f'=SUMIFS({M},{V},"Take")')
    R.cell(row=row, column=4, value=f"=C{row}/{HOURS}")
    R.cell(row=row, column=5, value=f"=C{row}-$C$5")
    R.cell(row=row, column=6, value=f'=COUNTIFS({J},"Took",{V},"<>Take")')
    R.cell(row=row, column=7, value=f'=COUNTIFS({J},"Passed",{V},"Take")')
last = 5 + len(rule_names)
for row in range(5, last + 1):
    for col in range(1, 8):
        c = R.cell(row=row, column=col)
        c.font = f_link if col in (2, 3, 6, 7) else f_base
        c.border = Border(bottom=thin)
    R.cell(row=row, column=3).number_format = MONEY
    R.cell(row=row, column=4).number_format = MONEY
    R.cell(row=row, column=5).number_format = '+$#,##0.00;-$#,##0.00;"-"'
R.cell(row=5, column=1).font = f_bold
R.conditional_formatting.add(f"E6:E{last}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FONT, size=10, color="1F6B4F", bold=True)))
R.conditional_formatting.add(f"E6:E{last}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, size=10, color="A3245E", bold=True)))
notes = [
    "Rides driven keep their real fare, tip, time and costs. Offers passed on are estimated from their offer cards (see Assumptions).",
    "Hours on the road are the same for every rule, so the per-hour column compares rules fairly.",
    "Say a rule in the app (\"What if I skip pickups over 4 miles?\") to replay one that isn't listed here.",
]
for i, t in enumerate(notes):
    R.cell(row=last + 2 + i, column=1, value=t).font = f_note
R.column_dimensions["A"].width = 38
for col, w in zip("BCDEFG", (9, 15, 14, 15, 15, 15)):
    R.column_dimensions[col].width = w
R.freeze_panes = "A5"

wb.calculation.fullCalcOnLoad = True  # Excel and Numbers compute the formulas when the file opens
out = HERE / "ride_offers_october.xlsx"
wb.save(out)
print("saved", out, N, "offers")
